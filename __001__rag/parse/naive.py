"""NaiveParser：本地轻量解析器（无 GPU / 无外部解析服务时的教学与兜底方案）。

用途：
  - 讲义"最小闭环"默认解析器（注册后即可本地跑通入库→检索→问答）；
  - 生产干净 PDF / DOCX / 文本文件的快速兜底。

支持：pdf（PyMuPDF）、docx（python-docx）、txt/md 纯文本。
复杂文档（扫描件/表格/公式）请使用 mineru / ocrflux 解析服务。

注册（ChunkMethod 表，讲义 02 章有完整 SQL）：
  method='naive', method_type='parser', method_class='__001__rag.parse.naive.NaiveParser'
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from loguru import logger

from __001__rag.parse.base import BaseParser, ParsedDocument, ParsedElement, ParserContext

_TEXT_SUFFIXES = {".txt", ".md", ".markdown"}
_DOCX_SUFFIXES = {".docx"}
_PDF_SUFFIXES = {".pdf"}


class NaiveParser(BaseParser):
    """本地文本提取解析器：文件字节 → Markdown 风格文本。"""

    parser_id = "naive"

    def supports(self, filename: str, content_type: str | None = None) -> bool:
        suffix = Path(filename or "").suffix.lower()
        return suffix in _TEXT_SUFFIXES | _DOCX_SUFFIXES | _PDF_SUFFIXES

    def parse(self, file_bytes: bytes, context: ParserContext) -> ParsedDocument:
        filename = context.filename or "untitled"
        suffix = Path(filename).suffix.lower()
        text = ""

        if suffix in _TEXT_SUFFIXES:
            text = self._parse_text(file_bytes)
        elif suffix in _DOCX_SUFFIXES:
            text = self._parse_docx(file_bytes)
        elif suffix in _PDF_SUFFIXES:
            text = self._parse_pdf(file_bytes)
        else:
            logger.info("naive 解析器跳过不支持的文件 filename={}", filename)
            text = ""

        elements: list[ParsedElement] = []
        if text:
            elements.append(
                ParsedElement(
                    content=text,
                    element_type="paragraph",
                    metadata={"parse_method": "naive"},
                )
            )
        logger.info(
            "naive 解析完成 filename={} chars={} elements={}",
            filename, len(text), len(elements),
        )
        return ParsedDocument(
            doc_id=context.doc_id,
            filename=filename,
            elements=elements,
            metadata={"parse_method": "naive"},
        )

    @staticmethod
    def _parse_text(file_bytes: bytes) -> str:
        for encoding in ("utf-8", "gbk", "latin-1"):
            try:
                return file_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        return file_bytes.decode("utf-8", errors="ignore")

    @staticmethod
    def _parse_docx(file_bytes: bytes) -> str:
        try:
            import io

            from docx import Document

            doc = Document(io.BytesIO(file_bytes))
            return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as exc:  # noqa: BLE001
            logger.exception("naive docx 解析失败 err={}", exc)
            return ""

    @staticmethod
    def _parse_pdf(file_bytes: bytes) -> str:
        try:
            import fitz  # PyMuPDF

            doc = fitz.open(stream=file_bytes, filetype="pdf")
            pages = [page.get_text("text") for page in doc]
            return "\n\n".join(pages)
        except Exception as exc:  # noqa: BLE001
            logger.exception("naive pdf 解析失败 err={}", exc)
            return ""


def _demo_supports() -> dict[str, bool]:
    """供讲义最小 demo 展示 supports 判定。"""
    parser = NaiveParser()
    return {
        "a.pdf": parser.supports("a.pdf"),
        "b.docx": parser.supports("b.docx"),
        "c.txt": parser.supports("c.txt"),
        "d.png": parser.supports("d.png"),
    }


if __name__ == "__main__":  # pragma: no cover
    from pprint import pprint

    pprint(_demo_supports())
